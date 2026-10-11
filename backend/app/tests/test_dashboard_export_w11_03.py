"""
W11-03 — Dashboard Excel export.

The export must reproduce what the dashboard shows: the same six KPI figures,
the same role scoping, and the same Lab Group / Lab dropdown narrowing.
"""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.api.dashboard import (
    SpendingSummary,
    _dashboard_kpis,
    _dashboard_range_label,
    _entity_label,
    _filter_breakdown,
)
from app.api.dashboard import BreakdownCategory, BreakdownLab, BreakdownLabGroup
from app.services.dashboard_report_service import (
    DashboardKpis,
    DashboardMeta,
    build_dashboard_workbook,
)


# ─── Builders ────────────────────────────────────────────────────────────
def _category(name: str = "Equipment", budget: str = "1000") -> BreakdownCategory:
    return BreakdownCategory(
        category_id=1,
        category_name=name,
        budget_sgd=budget,
        actual_spend_sgd="400",
        pending_sgd="100",
        available_balance_sgd="500",
        utilisation_pct=40.0,
        projected_utilisation_pct=50.0,
    )


def _lab(
    lab_id: int = 1, name: str = "W3 Food Lab", budget: str = "1000"
) -> BreakdownLab:
    return BreakdownLab(
        lab_id=lab_id,
        lab_name=name,
        budget_sgd=budget,
        actual_spend_sgd="400",
        pending_sgd="100",
        available_balance_sgd="500",
        utilisation_pct=40.0,
        projected_utilisation_pct=50.0,
        categories=[_category()],
    )


def _group(
    group_id: int = 1,
    name: str = "FCB",
    labs: list[BreakdownLab] | None = None,
) -> BreakdownLabGroup:
    return BreakdownLabGroup(
        lab_group_id=group_id,
        lab_group_name=name,
        budget_sgd="1000",
        actual_spend_sgd="400",
        pending_sgd="100",
        available_balance_sgd="500",
        utilisation_pct=40.0,
        projected_utilisation_pct=50.0,
        labs=labs if labs is not None else [_lab()],
    )


def _meta() -> DashboardMeta:
    return DashboardMeta(
        generated_at=datetime(2026, 9, 22, 15, 4),
        generated_by_name="Test Admin",
        generated_by_email="admin@test.com",
        role_label="Admin",
        scope_label="All Lab Groups (Admin — full access)",
        fy_label="FY2026 (01 Apr 2026 – 31 Mar 2027)",
        date_range_label="Whole year",
        lab_group_label="All Lab Groups",
        lab_label="All labs",
    )


def _kpis() -> DashboardKpis:
    return DashboardKpis(
        total_budget=1051000.0,
        actual_spend=705610.0,
        pending=0.0,
        available_balance=345390.0,
        utilisation_pct=67.1,
        projected_utilisation_pct=67.1,
    )


def _load(groups, monthly=()):
    stream: BytesIO = build_dashboard_workbook(
        _meta(), _kpis(), groups, groups, monthly
    )
    return load_workbook(stream)


def _find_row(ws, label: str) -> int:
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == label:
            return row
    raise AssertionError(f"row {label!r} not found")


# ─── KPI derivation (must mirror DashboardPage.tsx) ─────────────────────
def _spending(total="705610", pending="0", by_group=None, pending_by_group=None):
    return SpendingSummary(
        total_actual_spend_sgd=total,
        by_lab_group=by_group or {},
        total_pending_sgd=pending,
        pending_by_lab_group=pending_by_group or {},
    )


def test_kpis_sum_all_group_budgets_when_unfiltered() -> None:
    lab_groups = [{"id": 1, "budget": 600000.0}, {"id": 2, "budget": 451000.0}]
    kpis = _dashboard_kpis(_spending(), lab_groups, [], None, None)
    assert kpis.total_budget == pytest.approx(1051000.0)
    assert kpis.actual_spend == pytest.approx(705610.0)
    assert kpis.available_balance == pytest.approx(345390.0)
    assert kpis.utilisation_pct == pytest.approx(67.1, abs=0.1)


def test_kpis_read_lab_groups_as_dicts_not_attributes() -> None:
    """
    get_lab_groups() returns dicts; a defaulted getattr would silently report
    every budget as zero. Guard against that regression.
    """
    kpis = _dashboard_kpis(_spending(), [{"id": 1, "budget": 250000.0}], [], None, None)
    assert kpis.total_budget == pytest.approx(250000.0)


def test_kpis_scope_to_the_selected_lab_group() -> None:
    lab_groups = [{"id": 1, "budget": 600000.0}, {"id": 2, "budget": 451000.0}]
    spending = _spending(by_group={"2": "100000"}, pending_by_group={"2": "5000"})
    kpis = _dashboard_kpis(spending, lab_groups, [], 2, None)
    assert kpis.total_budget == pytest.approx(451000.0)
    assert kpis.actual_spend == pytest.approx(100000.0)
    assert kpis.pending == pytest.approx(5000.0)
    assert kpis.available_balance == pytest.approx(346000.0)


def test_kpis_read_straight_off_the_row_when_a_lab_is_selected() -> None:
    groups = [_group(labs=[_lab(lab_id=7, budget="8000")])]
    kpis = _dashboard_kpis(
        _spending(), [{"id": 1, "budget": 999999.0}], groups, None, 7
    )
    assert kpis.total_budget == pytest.approx(8000.0)
    assert kpis.actual_spend == pytest.approx(400.0)


def test_zero_budget_does_not_divide_by_zero() -> None:
    kpis = _dashboard_kpis(
        _spending(total="500"), [{"id": 1, "budget": 0.0}], [], None, None
    )
    assert kpis.utilisation_pct == 0.0
    assert kpis.projected_utilisation_pct == 0.0


# ─── Dropdown narrowing ──────────────────────────────────────────────────
def test_lab_group_filter_narrows_to_one_group() -> None:
    groups = [_group(1, "FCB"), _group(2, "ENG")]
    assert [g.lab_group_name for g in _filter_breakdown(groups, 2, None)] == ["ENG"]


def test_lab_filter_narrows_to_one_lab_inside_its_group() -> None:
    groups = [
        _group(1, "FCB", labs=[_lab(1, "Lab A"), _lab(2, "Lab B")]),
        _group(2, "ENG", labs=[_lab(3, "Lab C")]),
    ]
    result = _filter_breakdown(groups, None, 2)
    assert len(result) == 1
    assert [lab.lab_name for lab in result[0].labs] == ["Lab B"]


def test_no_filters_returns_everything() -> None:
    groups = [_group(1), _group(2)]
    assert len(_filter_breakdown(groups, None, None)) == 2


def test_filtering_does_not_mutate_the_source_tree() -> None:
    groups = [_group(1, "FCB", labs=[_lab(1, "Lab A"), _lab(2, "Lab B")])]
    _filter_breakdown(groups, None, 2)
    assert len(groups[0].labs) == 2


# ─── Workbook shape ──────────────────────────────────────────────────────
def test_workbook_has_three_sheets() -> None:
    wb = _load([_group()])
    assert wb.sheetnames == ["Dashboard", "Budget Breakdown", "Monthly Trend"]


def test_kpi_values_land_on_the_overview_sheet() -> None:
    ws = _load([_group()])["Dashboard"]
    assert ws.cell(row=_find_row(ws, "Total Budget"), column=2).value == pytest.approx(
        1051000.0
    )
    assert ws.cell(row=_find_row(ws, "Utilisation %"), column=2).value == pytest.approx(
        67.1
    )


def test_overview_carries_a_native_excel_chart() -> None:
    """A real chart object, not a pasted image, so it stays editable."""
    ws = _load([_group()])["Dashboard"]
    assert len(ws._charts) == 1


def test_breakdown_sheet_is_one_row_per_level() -> None:
    groups = [_group(1, "FCB", labs=[_lab(1, "Lab A")])]
    ws = _load(groups)["Budget Breakdown"]
    header = _find_row(ws, "Level")
    levels = [ws.cell(row=header + i, column=1).value for i in range(1, 4)]
    assert levels == ["Lab Group", "Lab", "Category"]


def test_breakdown_amounts_are_numeric() -> None:
    ws = _load([_group()])["Budget Breakdown"]
    cell = ws.cell(row=_find_row(ws, "Level") + 1, column=5)
    assert isinstance(cell.value, (int, float))
    assert cell.number_format == "#,##0.00"


def test_empty_breakdown_still_produces_a_valid_workbook() -> None:
    wb = _load([])
    ws = wb["Budget Breakdown"]
    assert "No budget data matched" in str(
        ws.cell(row=_find_row(ws, "Level") + 1, column=1).value
    )


class _Point:
    def __init__(self, month: str, amount: str) -> None:
        self.month = month
        self.amount = amount


def test_monthly_sheet_charts_the_points() -> None:
    wb = _load([_group()], monthly=[_Point("2026-04", "100"), _Point("2026-05", "250")])
    ws = wb["Monthly Trend"]
    header = _find_row(ws, "Month")
    assert ws.cell(row=header + 1, column=1).value == "2026-04"
    assert ws.cell(row=header + 2, column=2).value == pytest.approx(250.0)
    assert len(ws._charts) == 1


def test_monthly_sheet_handles_no_data() -> None:
    ws = _load([_group()], monthly=[])["Monthly Trend"]
    assert "No approved spending" in str(
        ws.cell(row=_find_row(ws, "Month") + 1, column=1).value
    )


def test_cover_block_records_scope_and_filters() -> None:
    ws = _load([_group()])["Dashboard"]
    text = "\n".join(
        str(ws.cell(row=r, column=c).value) for r in range(1, 13) for c in (1, 2)
    )
    assert "LAB BUDGET MONITORING SYSTEM" in text
    assert "admin@test.com" in text
    assert "All Lab Groups (Admin — full access)" in text


# ─── Cover labels ────────────────────────────────────────────────────────
def test_range_label_prefers_a_custom_window() -> None:
    assert _dashboard_range_label("90d", "2026-01-01", "2026-02-05") == (
        "Custom: 2026-01-01 to 2026-02-05"
    )
    assert _dashboard_range_label("30d", "", "") == "Last 30 days"


def test_entity_label_falls_back_when_unfiltered() -> None:
    assert _entity_label(None, {1: "FCB"}, "All Lab Groups") == "All Lab Groups"
    assert _entity_label(1, {1: "FCB"}, "All Lab Groups") == "FCB"


# ─── Role scoping through the HTTP layer ────────────────────────────────
@pytest.fixture
def export_client(db):
    from unittest.mock import MagicMock

    from fastapi.testclient import TestClient

    from app.database import get_db
    from app.dependencies.auth import get_current_user
    from app.main import app
    from app.models.auth import User
    from app.models.config import LabIcAssignment

    state: dict[str, User] = {}

    def _as(role: str, lab_group_id=None, lab_ids=()) -> None:
        state["user"] = User(
            user_id=999,
            email=f"{role}@test.com",
            full_name=f"Test {role}",
            role_id=1,
            role=MagicMock(role_name=role),
            lab_group_id=lab_group_id,
            lab_assignments=[LabIcAssignment(lab_id=lab_id) for lab_id in lab_ids],
        )

    previous = dict(app.dependency_overrides)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: state["user"]
    client = TestClient(app)
    client.login_as = _as  # type: ignore[attr-defined]
    try:
        yield client
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous)


def test_export_returns_an_xlsx_download(db, export_client) -> None:
    export_client.login_as("admin")
    response = export_client.get("/api/dashboard/export")
    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert ".xlsx" in response.headers["content-disposition"]
    load_workbook(BytesIO(response.content))  # parses cleanly


def test_export_rejects_an_unknown_role(db, export_client) -> None:
    export_client.login_as("vendor")
    assert export_client.get("/api/dashboard/export").status_code == 403


def test_teamlead_export_is_scoped_to_their_group(db, export_client) -> None:
    from app.models.config import LabGroup

    groups = (
        db.query(LabGroup).filter(LabGroup.lab_group_name != "General").limit(2).all()
    )
    if len(groups) < 2:
        pytest.skip("Need two seeded lab groups; run dev_seed.sql.")
    mine, theirs = groups[0], groups[1]

    export_client.login_as("teamlead", lab_group_id=mine.lab_group_id)
    wb = load_workbook(BytesIO(export_client.get("/api/dashboard/export").content))
    text = "\n".join(
        str(wb["Budget Breakdown"].cell(row=r, column=c).value)
        for r in range(1, wb["Budget Breakdown"].max_row + 1)
        for c in (2,)
    )
    assert theirs.lab_group_name not in text
    assert "Team Lead scope" in str(wb["Dashboard"].cell(row=7, column=2).value)


def test_export_honours_the_date_window(db, export_client) -> None:
    """A window in the far past leaves no monthly points."""
    export_client.login_as("admin")
    response = export_client.get(
        "/api/dashboard/export"
        f"?start_date={date(1990, 1, 1)}&end_date={date(1990, 12, 31)}"
    )
    ws = load_workbook(BytesIO(response.content))["Monthly Trend"]
    assert "No approved spending" in str(
        ws.cell(row=_find_row(ws, "Month") + 1, column=1).value
    )
